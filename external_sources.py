"""Read-only multi-format GitHub / NJU Box WebDAV adapters and session-local logical tables."""
import csv
import io
import json
from pathlib import PurePosixPath
import math
import re
import threading
import time
import zipfile
from datetime import datetime, timezone
from urllib.parse import quote, unquote, urlsplit
from xml.etree import ElementTree as ET

DEFAULT_GITHUB = 'https://raw.githubusercontent.com/neondatabase/postgres-sample-dbs/main/titanic.csv'
BOX_BASE = 'https://box.nju.edu.cn/seafdav/'
DEFAULT_EXCEL = '2021CESAI报告数据汇总/省市、部委数据（微信、双微指数）(1).xlsx'
MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 50000
MAX_CELLS = 500000

SUPPORTED_FORMATS = {'csv', 'xls', 'xlsx', 'json', 'jsonl', 'ndjson', 'tsv', 'txt', 'parquet'}

class SourceError(ValueError):
    pass


def dependencies():
    try:
        import duckdb
        import requests
        import openpyxl
        import certifi
        return duckdb, requests, openpyxl, certifi
    except ImportError:
        raise SourceError('请先安装 requirements-external.txt 中的依赖，然后重启平台。') from None


def quote_id(value):
    return '"' + value.replace('"', '""') + '"'


def safe_column_names(header, width):
    names, used = [], set()
    for i in range(width):
        name = str(header[i]).strip() if i < len(header) and header[i] is not None else ''
        base = name or f'column_{i+1}'
        name, n = base, 2
        while name.casefold() in used:
            name, n = f'{base}_{n}', n + 1
        used.add(name.casefold())
        names.append(name)
    return names


def github_url(value):
    u = urlsplit(value.strip())
    if u.scheme != 'https' or u.username or u.password or u.port or u.query or u.fragment:
        raise SourceError('请输入不含查询参数的 GitHub HTTPS 文件地址。')
    parts = unquote(u.path).strip('/').split('/')
    if u.hostname == 'github.com' and len(parts) >= 5 and parts[2] == 'blob':
        parts.pop(2)
    elif u.hostname != 'raw.githubusercontent.com':
        raise SourceError('仅支持 github.com 文件地址或 raw.githubusercontent.com。')
    if len(parts) < 4 or any(p in ('', '.', '..') for p in parts) or PurePosixPath(parts[-1]).suffix.lower().lstrip('.') not in SUPPORTED_FORMATS | {'rar', 'zip'}:
        raise SourceError('请指定支持格式的具体文件：CSV、XLS、XLSX、JSON、JSONL、NDJSON、TSV、TXT、Parquet、RAR、ZIP。')
    return 'https://raw.githubusercontent.com/' + quote('/'.join(parts), safe='/')


def box_url(path):
    path = path.strip().strip('/')
    if any(p in ('.', '..') for p in path.split('/')) or '\\' in path or '\x00' in path or '%' in path:
        raise SourceError('路径不能包含上级目录、反斜杠或百分号编码。请填写原始中文路径。')
    return BOX_BASE + quote(path, safe='/')


def fetch(url, method='GET', username='', token='', body=None):
    _, requests, _, certifi = dependencies()
    if url.startswith(BOX_BASE) and (not username.strip() or not token):
        raise SourceError('请输入南大云盘 WebDAV 用户名和令牌。')
    headers = {'Accept': 'application/octet-stream'}
    if method == 'PROPFIND':
        headers.update({'Depth': '1', 'Content-Type': 'application/xml'})
    auth = (username.strip(), token) if url.startswith(BOX_BASE) else None
    try:
        with requests.request(method, url, auth=auth, headers=headers, data=body,
                              verify=certifi.where(), timeout=(10, 40),
                              allow_redirects=False, stream=True) as response:
            expected = 207 if method == 'PROPFIND' else 200
            if response.status_code != expected:
                hints = {401:'认证失败，请检查 WebDAV 用户名和令牌', 403:'无访问权限或源端限制', 404:'文件或目录不存在'}
                raise SourceError(hints.get(response.status_code, f'源端返回 HTTP {response.status_code}；不跟随重定向'))
            chunks, size = [], 0
            started = time.monotonic()
            for chunk in response.iter_content(65536):
                size += len(chunk)
                if size > MAX_BYTES or time.monotonic() - started > 60:
                    raise SourceError('远程响应超过 20 MB 或读取超时。')
                chunks.append(chunk)
            return b''.join(chunks)
    except requests.RequestException:
        raise SourceError('远程连接失败，请检查网络、证书与源端状态。') from None


def list_directory(path, username, token):
    url = box_url(path).rstrip('/') + '/'
    payload = fetch(url, 'PROPFIND', username, token,
                    b'<d:propfind xmlns:d="DAV:"><d:prop><d:resourcetype/><d:getcontentlength/></d:prop></d:propfind>')
    if b'<!DOCTYPE' in payload.upper() or b'<!ENTITY' in payload.upper():
        raise SourceError('目录响应包含不支持的 XML 声明。')
    try:
        root = ET.fromstring(payload)
    except ET.ParseError:
        raise SourceError('云盘目录响应格式错误。') from None
    ns, items = {'d':'DAV:'}, []
    current = unquote(urlsplit(url).path).rstrip('/')
    for entry in root.findall('d:response', ns):
        href = urlsplit(entry.findtext('d:href', '', ns))
        if href.netloc and href.hostname != 'box.nju.edu.cn':
            continue
        item_path = unquote(href.path).rstrip('/')
        if item_path == current or not item_path.startswith(current + '/'):
            continue
        rel = item_path[len('/seafdav/'):]
        if '/' in item_path[len(current)+1:]:
            continue
        items.append({'name':item_path.rsplit('/',1)[-1], 'path':rel,
                      'directory':entry.find('.//d:collection', ns) is not None})
        if len(items) >= 1000:
            break
    return {'path':path, 'items':items, 'limit':1000}


def parse_csv(payload, model):
    if model not in ('raw', 'passengers'):
        raise SourceError('未知 CSV 映射模式。')
    try:
        reader = csv.reader(io.StringIO(payload.decode('utf-8-sig'), newline=''))
        header = next(reader)
        if not header or len(header) > 256 or len(header) != len(set(x.casefold() for x in header)):
            raise SourceError('CSV 表头为空、重复或超过 256 列。')
        rows = []
        for row in reader:
            if not row:
                continue
            if len(row) != len(header):
                raise SourceError('CSV 行与表头列数不一致。')
            rows.append(row)
            if len(rows) > MAX_ROWS or len(rows)*len(header) > MAX_CELLS:
                raise SourceError('数据超出 50000 行或 500000 单元格限制。')
    except (UnicodeError, csv.Error, StopIteration):
        raise SourceError('请提供有效的 UTF-8 CSV 文件。') from None
    if model == 'raw':
        return [('github_csv', 'CSV 原始字段', safe_column_names(header, len(header)), ['VARCHAR']*len(header), rows)]
    fields = ['passengerid','name','sex','age','survived','pclass','fare']
    if not set(fields).issubset(header):
        raise SourceError('乘客映射需要字段：' + ', '.join(fields) + '；其他 CSV 请选择原始字段模式。')
    indexes = [header.index(f) for f in fields]
    def number(v):
        if not v.strip(): return None
        n = float(v)
        if not math.isfinite(n): raise ValueError()
        return n
    mapped = []
    try:
        for row in rows:
            pid, name, sex, age, survived, pclass, fare = [row[i] for i in indexes]
            if survived.strip() not in ('','0','1','0.0','1.0'): raise ValueError()
            mapped.append((int(pid),name,sex,number(age),None if not survived.strip() else bool(int(float(survived))),int(pclass),number(fare)))
    except (ValueError, OverflowError):
        raise SourceError('CSV 中有无法转换的编号、数值或存活标志。') from None
    return [('passengers','乘客统一逻辑表',['passenger_id','name','sex','age','survived','passenger_class','fare'],
             ['INTEGER','VARCHAR','VARCHAR','DOUBLE','BOOLEAN','INTEGER','DOUBLE'],mapped)]


def parse_excel(payload, selected_sheets=None):
    _, _, openpyxl, _ = dependencies()
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as z:
            if sum(i.file_size for i in z.infolist()) > 80*1024*1024:
                raise SourceError('Excel 解压后超过 80 MB。')
        book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except SourceError: raise
    except Exception:
        raise SourceError('源端没有返回有效的 XLSX 文件。') from None
    tables, cells = [], 0
    try:
        if len(book.worksheets) > 20: raise SourceError('最多支持 20 个工作表。')
        for idx, sheet in enumerate(book.worksheets,1):
            if selected_sheets is not None and sheet.title not in selected_sheets: continue
            if (sheet.max_column or 0) > 256 or (sheet.max_row or 0) > MAX_ROWS+1:
                raise SourceError('工作表范围超过 256 列或 50000 数据行。')
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells += len(row)
                if cells > MAX_CELLS: raise SourceError('工作簿超过 500000 单元格。')
                if any(v is not None and str(v).strip() for v in row): rows.append(row)
            if not rows: continue
            width = max(map(len,rows))
            columns = safe_column_names(rows[0],width)
            values = [[None if i>=len(r) or r[i] is None else str(r[i]) for i in range(width)] for r in rows[1:]]
            tables.append((f'box_sheet_{idx}',sheet.title,columns,['VARCHAR']*width,values))
    finally: book.close()
    if not tables: raise SourceError('工作簿没有非空工作表。')
    return tables



def parse_xls(payload, selected_sheets=None):
    try:
        import xlrd
    except ImportError:
        raise SourceError('读取旧版 XLS 需要安装 requirements-external.txt 中的 xlrd。') from None
    book = None
    try:
        book = xlrd.open_workbook(file_contents=payload, on_demand=True)
        if book.nsheets>20: raise SourceError('最多支持 20 个工作表。')
        tables, cells = [], 0
        for index in range(book.nsheets):
            if selected_sheets is not None and book.sheet_names()[index] not in selected_sheets: continue
            sheet = book.sheet_by_index(index)
            if sheet.ncols>256 or sheet.nrows>MAX_ROWS+1:
                raise SourceError('工作表范围超过 256 列或 50000 数据行。')
            cells += sheet.nrows*sheet.ncols
            if cells>MAX_CELLS: raise SourceError('工作簿超过 500000 单元格。')
            rows = []
            for row_index in range(sheet.nrows):
                values = []
                for cell in sheet.row(row_index):
                    if cell.ctype in (xlrd.XL_CELL_EMPTY,xlrd.XL_CELL_BLANK): value=None
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        value=xlrd.xldate_as_datetime(cell.value,book.datemode).isoformat(sep=' ')
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN: value=str(bool(cell.value))
                    elif cell.ctype == xlrd.XL_CELL_ERROR: value=xlrd.error_text_from_code.get(cell.value,'#ERROR')
                    elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer(): value=str(int(cell.value))
                    else: value=str(cell.value)
                    values.append(value)
                if any(v is not None and v.strip() for v in values): rows.append(values)
            if rows:
                columns=safe_column_names(rows[0],sheet.ncols)
                tables.append((f'box_sheet_{index+1}',sheet.name,columns,['VARCHAR']*len(columns),rows[1:]))
            book.unload_sheet(index)
        if not tables: raise SourceError('工作簿没有非空工作表。')
        return tables
    except SourceError: raise
    except Exception:
        raise SourceError('无法解析 XLS：文件可能已损坏、加密，或并非真正的旧版 Excel 文件。') from None
    finally:
        if book is not None: book.release_resources()


def text_payload(payload):
    try:
        return payload.decode('utf-8-sig')
    except UnicodeError:
        raise SourceError('文本文件必须使用 UTF-8 编码。') from None


def records_table(records):
    if not records or not all(isinstance(r, dict) for r in records):
        raise SourceError('JSON 必须为非空对象数组、单个对象或逐行对象。')
    keys = list(dict.fromkeys(k for r in records for k in r))
    if not keys or len(keys) > 256 or len(records) > MAX_ROWS or len(keys)*len(records) > MAX_CELLS:
        raise SourceError('数据为空或超过 256 列、50000 行、500000 单元格限制。')
    def cell(v):
        if v is None: return None
        if isinstance(v, (dict, list, bool)):
            return json.dumps(v, ensure_ascii=False, allow_nan=False)
        return str(v)
    return safe_column_names(keys,len(keys)), [[cell(r.get(k)) for k in keys] for r in records]


def file_extension(location):
    path = unquote(urlsplit(location).path) if location.startswith(('https://','http://')) else location
    return PurePosixPath(path).suffix.lower().lstrip('.')


def parse_file(payload, location, source, model='raw', selected_sheets=None):
    """Dispatch by suffix; isolate table names by source, preserving old names."""
    if source not in ('github', 'box'): raise SourceError('未知数据来源。')
    ext = file_extension(location)
    if ext not in SUPPORTED_FORMATS: raise SourceError('不支持此文件格式。支持：' + ', '.join(sorted(SUPPORTED_FORMATS)))
    if len(payload) > MAX_BYTES: raise SourceError('文件超过 20 MB。')
    if ext in ('xls','xlsx'):
        return [(name.replace('box_',source+'_',1),label,cols,types,rows)
                for name,label,cols,types,rows in (parse_xls(payload,selected_sheets) if ext=='xls' else parse_excel(payload,selected_sheets))]
    if ext == 'csv':
        mode = model
        if mode == 'auto':
            try: header = next(csv.reader(io.StringIO(text_payload(payload))))
            except StopIteration: header = []
            mode = 'passengers' if {'passengerid','name','sex','age','survived','pclass','fare'}.issubset(header) and source=='github' else 'raw'
        tables = parse_csv(payload,mode)
        return [(name if source=='github' else 'box_csv',label,cols,types,rows) for name,label,cols,types,rows in tables]
    try:
        if ext in ('json','jsonl','ndjson'):
            def invalid_constant(value): raise ValueError('Non-finite JSON value')
            def decode(value): return json.loads(value, parse_constant=invalid_constant)
            if ext == 'json':
                records = decode(text_payload(payload))
                if isinstance(records,dict): records = [records]
            else:
                records = []
                for line in io.StringIO(text_payload(payload)):
                    if line.strip(): records.append(decode(line))
                    if len(records)>MAX_ROWS: raise SourceError('超过 50000 行限制。')
            columns, rows = records_table(records)
        elif ext == 'tsv':
            reader = csv.reader(io.StringIO(text_payload(payload),newline=''),delimiter='\t')
            header = next(reader)
            if not header or len(header)>256: raise SourceError('TSV 表头为空或超过 256 列。')
            columns, rows = safe_column_names(header,len(header)), []
            for row in reader:
                if not row: continue
                if len(row)!=len(columns): raise SourceError('TSV 行与表头列数不一致。')
                rows.append(row)
                if len(rows)>MAX_ROWS or len(rows)*len(columns)>MAX_CELLS: raise SourceError('TSV 数据超出行数或单元格限制。')
        elif ext == 'txt':
            columns, rows = ['line_number','content'], []
            for i,line in enumerate(io.StringIO(text_payload(payload)),1):
                if i>MAX_ROWS: raise SourceError('TXT 超过 50000 行。')
                rows.append([str(i),line.rstrip('\r\n')])
            if not rows: raise SourceError('TXT 文件为空。')
        else:
            try:
                import pyarrow.parquet as pq
            except ImportError:
                raise SourceError('Parquet 需要安装 requirements-external.txt 中的 pyarrow。') from None
            with pq.ParquetFile(io.BytesIO(payload)) as file:
                meta = file.metadata
                if meta.num_rows>MAX_ROWS or meta.num_columns>256 or meta.num_rows*meta.num_columns>MAX_CELLS:
                    raise SourceError('Parquet 超出行数、列数或单元格限制。')
                if sum(meta.row_group(i).total_byte_size for i in range(meta.num_row_groups))>80*1024*1024:
                    raise SourceError('Parquet 解压后超过 80 MB。')
                arrow = file.read(use_threads=False)
                columns = safe_column_names(arrow.column_names,len(arrow.column_names))
                if not columns: raise SourceError('Parquet 没有字段。')
                values = [c.to_pylist() for c in arrow.columns]
                rows = [[None if v is None else json.dumps(v,ensure_ascii=False,default=str) if isinstance(v,(dict,list)) else str(v)
                         for v in row] for row in zip(*values)]
        return [(source+'_'+ext, ext.upper()+' 文件', columns, ['VARCHAR']*len(columns), rows)]
    except SourceError: raise
    except Exception:
        raise SourceError('文件内容与格式不符，或无法解析；文本请使用 UTF-8，JSON 每条记录必须为对象。') from None


class LogicalSession:
    def __init__(self):
        duckdb, *_ = dependencies()
        self.con = duckdb.connect(':memory:',config={'enable_external_access':'false',
            'autoload_known_extensions':'false','autoinstall_known_extensions':'false',
            'threads':'1','memory_limit':'128MB','max_temp_directory_size':'0B'})
        self.lock = threading.Lock()
        self.scan_job = None
        self.cloud_credentials = None
        self.virtual_files = {}
        self.virtual_lock = threading.Lock()
        self.tables = {}
        self.events = []
        self.touched = time.monotonic()

    def close(self):
        with self.virtual_lock:
            self.virtual_files.clear()
            self.scan_job = None
            self.cloud_credentials = None
        with self.lock: self.con.close()

    def snapshot(self):
        with self.lock:
            return {'tables':list(self.tables.values()),'events':list(self.events),
                    'mode':'远程文件按需读取，内存快照查询；刷新后重新获取源文件'}

    def replace(self, source, location, tables):
        with self.lock:
            self.con.execute('BEGIN')
            metadata = {k:v for k,v in self.tables.items() if v['source']!=source}
            try:
                for name, info in self.tables.items():
                    if info['source']==source: self.con.execute('DROP TABLE '+quote_id(name))
                now = datetime.now(timezone.utc).isoformat()
                for name,label,columns,types,rows in tables:
                    schema = ', '.join(quote_id(c)+' '+t for c,t in zip(columns,types))
                    self.con.execute('CREATE TABLE '+quote_id(name)+' ('+schema+')')
                    if rows: self.con.executemany('INSERT INTO '+quote_id(name)+' VALUES ('+','.join('?' for _ in columns)+')', rows)
                    metadata[name]={'name':name,'label':label,'source':source,'location':location,
                                    'rows':len(rows),'columns':[{'name':c,'type':t} for c,t in zip(columns,types)],'refreshed_at':now}
                self.con.execute('COMMIT')
            except Exception:
                self.con.execute('ROLLBACK')
                raise SourceError('逻辑表载入失败，已保留上一份可用快照。') from None
            self.tables = metadata
            self.events.insert(0,{'time':now,'source':source,'message':f'刷新完成：{len(tables)} 张逻辑表'})
            self.events = self.events[:30]
        return self.snapshot()

    def query(self, sql, parameters=None):
        if not sql.strip() or len(sql)>10000: raise SourceError('请输入不超过 10000 字符的查询。')
        with self.lock:
            try:
                statements = self.con.extract_statements(sql)
                if len(statements)!=1 or str(statements[0].type)!='StatementType.SELECT':
                    raise SourceError('仅允许一条 SELECT 或 WITH 查询，不允许写入或修改配置。')
                timer = threading.Timer(5, self.con.interrupt)
                timer.start()
                started = time.perf_counter()
                try:
                    res = self.con.execute(statements[0], parameters or [])
                    columns = [c[0] for c in res.description]
                    rows = res.fetchmany(201)
                finally:
                    timer.cancel()
                    timer.join()
                # Arrays preserve duplicate output column names without silently overwriting values.
                return {'columns':columns,'rows':rows[:200],'truncated':len(rows)>200,
                        'limit':200,'elapsed_ms':round((time.perf_counter()-started)*1000,2)}
            except SourceError: raise
            except Exception:
                raise SourceError('查询失败：请检查表名、字段和 SQL；外部文件访问已禁用，执行最长 5 秒。') from None
